# -*- coding: utf-8 -*-
"""
Web App: AI Assistant - Full web UI + API
I. Doi soat du lieu (dien, nuoc uong)
II. Phan bo chi phi (dien, nuoc uong)
"""
import os
import re
import io
import tempfile
import shutil
from datetime import datetime
import calendar

from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import pdfplumber
import openpyxl

try:
    import pytesseract
    from PIL import Image
    HAS_OCR = True
except ImportError:
    HAS_OCR = False

# ==========================================
# CONFIG (same as desktop app)
# ==========================================
RULES = {
    "dien": {
        "name": "Phân bổ chi phí điện",
        "data_start_row": 5, "col_code": 8, "col_fill": 15, "sheet_index": 1,
        "excel_prefix": "Bang phan bo chi phi dien",
        "output_format": "Bảng phân bổ chi phí điện tháng {month} năm {year}",
        "title_format": "BẢNG PHÂN BỔ CHI PHÍ THÁNG {month}/{year}",
        "sheet_format": "Phan bo_{month}.26",
        "has_date_shift": True, "col_ky_tt": 13, "col_ky_tt_truoc": 14,
        "split_duplicates": False, "has_di_doi": True, "ref_month": 7,
    },
    "nuoc": {
        "name": "Phân bổ chi phí nước uống",
        "data_start_row": 5, "col_code": 2, "col_fill": 8, "sheet_index": 2,
        "excel_prefix": "Bang phan bo chi phi nuoc uong",
        "output_format": "Bảng phân bổ chi phí nước uống tháng {month} năm {year}",
        "title_format": "BẢNG PHÂN BỔ CHI PHÍ NƯỚC UỐNG THÁNG {month}/{year}",
        "sheet_format": "T{month}",
        "has_date_shift": False, "col_ky_tt": 13, "col_ky_tt_truoc": 14,
        "split_duplicates": True, "has_di_doi": False, "ref_month": 7,
    },
}

# ==========================================
# LOGIC (same as desktop app)
# ==========================================
def add_months(ds, n):
    try:
        d = datetime.strptime(ds.strip(), '%d/%m/%Y')
        t = d.year * 12 + d.month - 1 + n
        y, m = t // 12, t % 12 + 1
        return d.replace(year=y, month=m, day=min(d.day, calendar.monthrange(y, m)[1])).strftime('%d/%m/%Y')
    except: return None

def shift_dr(s, n):
    if not s: return None
    m = re.match(r'(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})', str(s).strip())
    if not m: return None
    a, b = add_months(m.group(1), n), add_months(m.group(2), n)
    return f"{a} - {b}" if a and b else None

def _code(t):
    if not t: return ""
    t = str(t).strip()
    m = re.match(r'^([A-Z]{2}[0-9A-Z]+)', t) or re.search(r'\b([A-Z]{2}\d{8,})\b', t)
    return m.group(1) if m else ""

def _amt(d, c):
    s = ""
    if d and d != '0': s = d.replace(',', '').replace('.00', '').strip()
    elif c and c != '0': s = c.replace(',', '').replace('.00', '').strip()
    return float(s) if s else 0

def extract_pdf_dien(fp):
    data = []
    with pdfplumber.open(fp) as pdf:
        for p in pdf.pages:
            for t in p.extract_tables():
                for r in t:
                    if not r or not r[0]: continue
                    if not str(r[0]).strip().isdigit(): continue
                    code = _code(str(r[5]) if r[5] else "")
                    amt = _amt(str(r[6]) if r[6] else "", str(r[7]) if r[7] else "")
                    if code and code != 'MSB' and amt > 0:
                        data.append({'code': code, 'amount': amt})
    return data

def extract_pdf_nuoc(fp):
    result = {}
    with pdfplumber.open(fp) as pdf:
        all_text = ""
        all_tables = []
        for page in pdf.pages:
            all_text += (page.extract_text() or "") + "\n"
            for table in page.extract_tables():
                all_tables.append(table)
        for table in all_tables:
            for row in table:
                if not row or not row[0]: continue
                m = re.match(r'^(C\d+)', str(row[0]).strip())
                if not m: continue
                code = m.group(1)
                qty_str = str(row[6]).strip() if len(row) > 6 and row[6] else ""
                try: qty = float(qty_str.replace(',', ''))
                except: qty = 0
                if qty > 0 and code not in result: result[code] = qty
        lines = all_text.split('\n')
        current_code = None
        for i, line in enumerate(lines):
            line = line.strip()
            m = re.match(r'^(C\d+)', line)
            if m: current_code = m.group(1)
            if ('Theo mat hang' in line or 'By SKU' in line) and current_code:
                if current_code not in result and i + 1 < len(lines):
                    sku_line = lines[i + 1].strip()
                    nums = re.findall(r'(\d[\d,]*)', sku_line)
                    if nums:
                        try:
                            qty = float(nums[-1].replace(',', ''))
                            if qty > 0: result[current_code] = qty
                        except: pass
    return [{'code': k, 'amount': v} for k, v in result.items()]

def extract_pdf_nuoc_by_pgh(fp):
    result = {}
    with pdfplumber.open(fp) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if 'PHIẾU GIAO HÀNG' in text or 'DELIVERY TICKET' in text:
                m = re.search(r'SỐ PHIẾU[/\s]*NO[\.\s]*\s*(\d{9})', text)
                if not m:
                    lines = text.split('\n')
                    for i, line in enumerate(lines):
                        if 'SỐ PHIẾU' in line or 'NO.' in line:
                            for j in range(i, min(i + 4, len(lines))):
                                nums = re.findall(r'\d{9}', lines[j])
                                if nums: m = re.match(r'(\d{9})', nums[0]); break
                            if m: break
                if m:
                    pgh = m.group(1)
                    tables = page.extract_tables()
                    for table in tables:
                        if len(table) >= 2:
                            header, row1 = table[0], table[1]
                            for ci, h in enumerate(header):
                                if h and 'lượng' in str(h).lower():
                                    try: qty = float(str(row1[ci]).replace(',', ''))
                                    except: qty = 0
                                    if pgh not in result: result[pgh] = qty
                                    break
                            break
            else:
                for table in page.extract_tables():
                    for row in table:
                        if not row: continue
                        col1 = str(row[1]).strip() if len(row) > 1 and row[1] else ""
                        col2 = str(row[2]).strip() if len(row) > 2 and row[2] else ""
                        col6 = str(row[6]).strip() if len(row) > 6 and row[6] else ""
                        if re.match(r'\d{2}/\d{2}/\d{4}', col1):
                            m = re.search(r'\d{9}', col2)
                            if m:
                                pgh = m.group(0)
                                try: qty = float(col6.replace(',', ''))
                                except: qty = 0
                                if pgh not in result: result[pgh] = qty
    return [{'code': k, 'amount': v} for k, v in result.items()]

def extract_xl_by_pgh(fp):
    result = {}
    wb = openpyxl.load_workbook(fp, keep_vba=False, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            if not row: continue
            pgh_val = row[4] if len(row) > 4 else None
            qty_val = row[15] if len(row) > 15 else None
            if pgh_val:
                m = re.search(r'\d{9}', str(pgh_val))
                if m and qty_val:
                    try: qty = float(str(qty_val).replace(',', ''))
                    except: qty = 0
                    if qty > 0 and m.group(0) not in result: result[m.group(0)] = qty
    wb.close()
    return [{'code': k, 'amount': v} for k, v in result.items()]

def extract_excel_data(fp):
    data = []
    wb = openpyxl.load_workbook(fp, keep_vba=False, data_only=True)
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5, values_only=True):
            if not r or not r[0]: continue
            code = str(r[0]).strip()
            amt = 0
            if len(r) > 2 and r[2]:
                try: amt = float(str(r[2]).replace(',', '').replace('.00', '').strip())
                except: pass
            if code and amt > 0: data.append({'code': code, 'amount': amt})
    wb.close()
    return data

def build_lk(data):
    lk = {}
    for i in data:
        if i['code'] not in lk: lk[i['code']] = i
    return lk

def process_phan_bo(ep, lk, rule_key, ms=0, year="2026"):
    r = RULES[rule_key]
    wb = openpyxl.load_workbook(ep, keep_vba=False)
    ws = wb.worksheets[r["sheet_index"]]
    res = []; u = sd = sn = 0
    merged = list(ws.merged_cells.ranges)
    cc = {}
    if r["split_duplicates"]:
        for ri in range(r["data_start_row"], ws.max_row + 1):
            c = ws.cell(row=ri, column=r["col_code"]).value
            if c: cc[str(c).strip()] = cc.get(str(c).strip(), 0) + 1
    for ri in range(r["data_start_row"], ws.max_row + 1):
        ec = ws.cell(row=ri, column=r["col_code"]).value
        if not ec or not str(ec).strip(): continue
        ec = str(ec).strip()
        if "HO HCM" in ec.upper(): continue
        cell = ws.cell(row=ri, column=r["col_fill"])
        if any(cell.coordinate in mr for mr in merged): continue
        if r["has_di_doi"]:
            c13 = ws.cell(row=ri, column=r["col_ky_tt"]).value or ""
            c14 = ws.cell(row=ri, column=r["col_ky_tt_truoc"]).value or ""
            if any(k in str(c13).lower() for k in ['di d', 'dời', 'ngừng']) or \
               any(k in str(c14).lower() for k in ['di d', 'dời', 'ngừng']):
                cell.value = None; sd += 1
                res.append({'row': ri, 'code': ec, 'action': 'SKIP', 'amount': 'blank'}); continue
            if ms and r["has_date_shift"]:
                n13 = shift_dr(c13, ms)
                if n13: ws.cell(row=ri, column=r["col_ky_tt"]).value = n13
                n14 = shift_dr(c14, ms)
                if n14: ws.cell(row=ri, column=r["col_ky_tt_truoc"]).value = n14
        if ec in lk:
            a = lk[ec]['amount']
            if r["split_duplicates"] and cc.get(ec, 1) > 1: a = a / cc[ec]
            cell.value = a; u += 1
            res.append({'row': ri, 'code': ec, 'action': 'UPDATED', 'amount': a})
        else:
            cell.value = 0; sn += 1
            res.append({'row': ri, 'code': ec, 'action': 'NO MATCH', 'amount': 0})
    mn = str(r["ref_month"] + ms).zfill(2)
    ws.cell(row=3, column=1).value = r["title_format"].format(month=mn, year=year)
    ws.title = r["sheet_format"].format(month=mn)
    return wb, res, u, sd, sn

def process_doi_soat(f1, f2, key):
    def extract(fp):
        ext = os.path.splitext(fp)[1].lower()
        if ext == '.pdf':
            return extract_pdf_nuoc_by_pgh(fp) if key == 'nuoc' else extract_pdf_dien(fp)
        elif ext in ('.xlsx', '.xls'): return extract_xl_by_pgh(fp)
        return []
    d1, d2 = extract(f1), extract(f2)
    lk1, lk2 = build_lk(d1), build_lk(d2)
    result = {}
    for pgh in set(list(lk1.keys()) + list(lk2.keys())):
        q1 = lk1.get(pgh, {}).get('amount', None)
        q2 = lk2.get(pgh, {}).get('amount', None)
        if q1 is not None and q2 is not None:
            result[pgh] = 'Đúng' if abs(q1 - q2) < 1 else 'Sai'
        else:
            result[pgh] = 'Không có kết quả'
    return result, len(d1), len(d2)

# ==========================================
# WEB APP
# ==========================================
app = FastAPI(title="AI Assistant", version="4.0")

HTML = """
<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>AI Assistant</title>
<style>
* { margin:0; padding:0; box-sizing:border-box; font-family:'Segoe UI',sans-serif; }
body { display:flex; height:100vh; background:#f5f5f5; }
.sidebar { width:240px; background:#263238; color:#fff; padding:0; overflow-y:auto; }
.sidebar h1 { background:#1a237e; padding:15px 20px; font-size:18px; }
.sidebar .team { background:#1a237e; padding:5px 20px 10px; font-size:11px; color:#64ffda; text-align:right; }
.sidebar .section { padding:10px 20px 5px; color:#18BC9C; font-size:12px; font-weight:bold; }
.sidebar button { display:block; width:100%; text-align:left; background:#263238; color:#b0bec5;
  border:none; padding:10px 20px; font-size:13px; cursor:pointer; }
.sidebar button:hover { background:#37474f; color:#fff; }
.sidebar button.active { background:#37474f; color:#fff; }
.sidebar .history { margin-top:20px; padding:10px 20px; color:#90a4ae; font-size:11px; border-top:1px solid #37474f; }
.content { flex:1; padding:30px; overflow-y:auto; }
.content h2 { color:#1a237e; font-size:22px; margin-bottom:5px; }
.content .desc { color:#666; font-size:13px; margin-bottom:20px; }
.step { background:#fff; border-radius:8px; padding:15px 20px; margin-bottom:15px; box-shadow:0 1px 3px rgba(0,0,0,0.1); }
.step h3 { color:#1a237e; font-size:14px; margin-bottom:10px; }
.step .num { display:inline-block; background:#1a237e; color:#fff; width:24px; height:24px;
  border-radius:50%; text-align:center; line-height:24px; font-size:12px; font-weight:bold; margin-right:8px; }
.file-input { border:2px dashed #ccc; border-radius:8px; padding:20px; text-align:center; cursor:pointer; }
.file-input:hover { border-color:#1a237e; }
.file-input input { display:none; }
.file-name { color:green; font-size:13px; margin-top:5px; }
.btn { background:#1a237e; color:#fff; border:none; padding:10px 25px; border-radius:5px;
  font-size:14px; cursor:pointer; margin-right:10px; }
.btn:hover { background:#283593; }
.btn:disabled { background:#999; cursor:not-allowed; }
.btn-green { background:#2e7d32; }
.btn-green:hover { background:#388e3c; }
table { width:100%; border-collapse:collapse; margin-top:10px; font-size:13px; }
th { background:#1a237e; color:#fff; padding:8px; text-align:left; }
td { padding:6px 8px; border-bottom:1px solid #eee; }
tr.updated { color:green; }
tr.mismatch { color:red; }
tr.skip { color:orange; }
.summary { background:#e3f2fd; padding:10px 15px; border-radius:5px; margin:10px 0; font-size:13px; }
.status { font-size:13px; margin-top:10px; }
.hidden { display:none; }
</style>
</head>
<body>
<div class="sidebar">
  <h1>AI Assistant</h1>
  <div class="team">MSB AI Hackathon 2026</div>
  <div class="section">I. ĐỐI SOÁT DỮ LIỆU</div>
  <button onclick="showView('ds-dien')" id="btn-ds-dien">I.1. Đối soát dữ liệu điện</button>
  <button onclick="showView('ds-nuoc')" id="btn-ds-nuoc">I.2. Đối soát dữ liệu nước uống</button>
  <div class="section">II. PHÂN BỔ CHI PHÍ</div>
  <button onclick="showView('pb-dien')" id="btn-pb-dien">II.1. Phân bổ chi phí điện</button>
  <button onclick="showView('pb-nuoc')" id="btn-pb-nuoc">II.2. Phân bổ chi phí nước uống</button>
  <div class="section">III. QUẢN LÝ CHI PHÍ</div>
  <button disabled style="color:#546e7a;font-style:italic">(Sắp có)</button>
  <div class="history" id="history">Chưa có lịch sử</div>
</div>
<div class="content" id="content"></div>
<script>
let currentView = '';
function showView(view) {
  currentView = view;
  document.querySelectorAll('.sidebar button').forEach(b => b.classList.remove('active'));
  const btn = document.getElementById('btn-' + view);
  if (btn) btn.classList.add('active');
  const isDS = view.startsWith('ds-');
  const isNuoc = view.endsWith('nuoc');
  const rule = isNuoc ? 'nuoc' : 'dien';
  const name = isDS ? (isNuoc ? 'Đối soát dữ liệu nước uống' : 'Đối soát dữ liệu điện')
                    : (isNuoc ? 'Phân bổ chi phí nước uống' : 'Phân bổ chi phí điện');
  const desc = isDS ? 'So sánh 2 file theo Số PGH (9 chữ số)'
                    : 'Cập nhật Excel template từ file đầu vào theo mã khách hàng';
  let html = `<h2>${name}</h2><div class="desc">${desc}</div>`;
  if (isDS) {
    html += `
    <div class="step"><h3><span class="num">1</span>Chọn file 1 (PDF/Excel)</h3>
      <div class="file-input" onclick="document.getElementById('f1').click()">
        📁 Chọn file 1 <input type="file" id="f1" onchange="showFile(this,'fn1')">
        <div class="file-name" id="fn1"></div></div></div>
    <div class="step"><h3><span class="num">2</span>Chọn file 2 (PDF/Excel)</h3>
      <div class="file-input" onclick="document.getElementById('f2').click()">
        📁 Chọn file 2 <input type="file" id="f2" onchange="showFile(this,'fn2')">
        <div class="file-name" id="fn2"></div></div></div>
    <div class="step"><h3><span class="num">3</span>Đối soát & Tải kết quả</h3>
      <button class="btn" onclick="doiSoat('${rule}')">Đối soát</button>
      <button class="btn btn-green" id="dlBtn" onclick="downloadDS('${rule}')" disabled>Tải Excel</button>
      <div class="status" id="status"></div></div>
    <div class="step"><h3><span class="num">4</span>Kết quả</h3>
      <div class="summary" id="summary">Chưa đối soát</div>
      <table id="results"><thead><tr><th>Số PGH</th><th>File 1</th><th>File 2</th><th>Kết quả</th></tr></thead>
      <tbody></tbody></table></div>`;
  } else {
    html += `
    <div class="step"><h3><span class="num">1</span>Chọn file đầu vào (PDF/Excel)</h3>
      <div class="file-input" onclick="document.getElementById('pf').click()">
        📁 Chọn file đầu vào <input type="file" id="pf" onchange="showFile(this,'pfn')">
        <div class="file-name" id="pfn"></div></div></div>
    <div class="step"><h3><span class="num">2</span>Chọn file Excel gốc (template)</h3>
      <div class="file-input" onclick="document.getElementById('tf').click()">
        📁 Chọn file Excel gốc <input type="file" id="tf" onchange="showFile(this,'tfn')">
        <div class="file-name" id="tfn"></div></div></div>
    <div class="step"><h3><span class="num">3</span>Xử lý & Tải kết quả</h3>
      <button class="btn" onclick="phanBo('${rule}')">Xử lý</button>
      <button class="btn btn-green" id="dlBtn" onclick="downloadPB('${rule}')" disabled>Tải Excel</button>
      <div class="status" id="status"></div></div>
    <div class="step"><h3><span class="num">4</span>Kết quả</h3>
      <div class="summary" id="summary">Chưa xử lý</div>
      <table id="results"><thead><tr><th>Dòng</th><th>Mã KH</th><th>Hành động</th><th>Giá trị</th></tr></thead>
      <tbody></tbody></table></div>`;
  }
  document.getElementById('content').innerHTML = html;
}
function showFile(input, nameId) {
  if (input.files[0]) document.getElementById(nameId).textContent = input.files[0].name;
}
let lastResult = null;
async function phanBo(rule) {
  const pf = document.getElementById('pf').files[0];
  const tf = document.getElementById('tf').files[0];
  if (!pf || !tf) { alert('Vui lòng chọn đủ 2 file!'); return; }
  document.getElementById('status').innerHTML = '<span style="color:blue">Đang xử lý...</span>';
  const fd = new FormData();
  fd.append('rule', rule); fd.append('month', '7');
  fd.append('input_file', pf); fd.append('excel_template', tf);
  try {
    const r = await fetch('/phan-bo', {method:'POST', body:fd});
    const d = await r.json();
    lastResult = {type:'pb', rule, pf, tf};
    document.getElementById('dlBtn').disabled = false;
    document.getElementById('status').innerHTML = '<span style="color:green">Xong!</span>';
    document.getElementById('summary').innerHTML = `Updated: ${d.updated} | Skip: ${d.skipped_di_doi} | No match: ${d.no_match} | Total: ${d.total}`;
    let rows = '';
    for (const x of d.results.slice(0,50)) {
      const cls = x.action==='UPDATED'?'updated':(x.action.includes('SKIP')?'skip':'');
      rows += `<tr class="${cls}"><td>${x.row}</td><td>${x.code}</td><td>${x.action}</td><td>${typeof x.amount==='number'?x.amount.toLocaleString():x.amount}</td></tr>`;
    }
    document.querySelector('#results tbody').innerHTML = rows;
  } catch(e) { document.getElementById('status').innerHTML = '<span style="color:red">Lỗi: '+e+'</span>'; }
}
async function doiSoat(rule) {
  const f1 = document.getElementById('f1').files[0];
  const f2 = document.getElementById('f2').files[0];
  if (!f1 || !f2) { alert('Vui lòng chọn 2 file!'); return; }
  document.getElementById('status').innerHTML = '<span style="color:blue">Đang đối soát...</span>';
  const fd = new FormData();
  fd.append('rule', rule); fd.append('file1', f1); fd.append('file2', f2);
  try {
    const r = await fetch('/doi-soat', {method:'POST', body:fd});
    const d = await r.json();
    lastResult = {type:'ds', rule, f1, f2};
    document.getElementById('dlBtn').disabled = false;
    document.getElementById('status').innerHTML = '<span style="color:green">Xong!</span>';
    document.getElementById('summary').innerHTML = `File1: ${d.file1_count} | File2: ${d.file2_count} | Đúng: ${d.matched} | Sai: ${d.mismatched} | Không có: ${d.no_result} | Total: ${d.total}`;
  } catch(e) { document.getElementById('status').innerHTML = '<span style="color:red">Lỗi: '+e+'</span>'; }
}
async function downloadPB(rule) {
  if (!lastResult || lastResult.type !== 'pb') return;
  const fd = new FormData();
  fd.append('rule', rule); fd.append('month', '7');
  fd.append('input_file', lastResult.pf); fd.append('excel_template', lastResult.tf);
  const r = await fetch('/phan-bo/download', {method:'POST', body:fd});
  const blob = await r.blob();
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a'); a.href = url;
  a.download = r.headers.get('content-disposition')?.split('filename=')[1]?.replace(/"/g,'') || 'output.xlsx';
  a.click(); URL.revokeObjectURL(url);
}
async function downloadDS(rule) {
  if (!lastResult || lastResult.type !== 'ds') return;
  const fd = new FormData();
  fd.append('rule', rule); fd.append('file1', lastResult.f1); fd.append('file2', lastResult.f2);
  const r = await fetch('/doi-soat/download', {method:'POST', body:fd});
  const blob = await r.blob();
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a'); a.href = url;
  a.download = 'Doi soat.xlsx'; a.click(); URL.revokeObjectURL(url);
}
showView('pb-dien');
</script>
</body></html>
"""

@app.get("/", response_class=HTMLResponse)
async def home():
    return HTML

@app.get("/health")
async def health():
    return {"status": "ok"}

@app.post("/phan-bo")
async def phan_bo(rule: str = Form(...), month: int = Form(...),
                  input_file: UploadFile = File(...), excel_template: UploadFile = File(...)):
    if rule not in RULES: raise HTTPException(400, f"Invalid rule: {rule}")
    tmpdir = tempfile.mkdtemp()
    try:
        ip = os.path.join(tmpdir, input_file.filename)
        with open(ip, 'wb') as f: f.write(await input_file.read())
        ep = os.path.join(tmpdir, excel_template.filename)
        with open(ep, 'wb') as f: f.write(await excel_template.read())
        ext = os.path.splitext(input_file.filename)[1].lower()
        if ext == '.pdf':
            data = extract_pdf_nuoc(ip) if rule == 'nuoc' else extract_pdf_dien(ip)
        elif ext in ('.xlsx', '.xls'): data = extract_excel_data(ip)
        else: raise HTTPException(400, f"Unsupported: {ext}")
        lk = build_lk(data)
        bn = os.path.basename(input_file.filename)
        m = re.search(r'(?:tháng|thang)\s*(\d+)[.\s]*(?:năm\s*)?(\d{4})?', bn, re.IGNORECASE)
        ms = int(m.group(1)) - RULES[rule]["ref_month"] if m else 0
        year = m.group(2) if m and m.group(2) else "2026"
        wb, res, u, sd, sn = process_phan_bo(ep, lk, rule, ms, year)
        return {"rule": rule, "updated": u, "skipped_di_doi": sd, "no_match": sn,
                "total": u+sd+sn, "results": [{**r, "amount": r["amount"] if isinstance(r["amount"], (int,float)) else str(r["amount"])} for r in res[:50]]}
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

@app.post("/phan-bo/download")
async def phan_bo_download(rule: str = Form(...), month: int = Form(...),
                           input_file: UploadFile = File(...), excel_template: UploadFile = File(...)):
    if rule not in RULES: raise HTTPException(400, f"Invalid rule: {rule}")
    tmpdir = tempfile.mkdtemp()
    try:
        ip = os.path.join(tmpdir, input_file.filename)
        with open(ip, 'wb') as f: f.write(await input_file.read())
        ep = os.path.join(tmpdir, excel_template.filename)
        with open(ep, 'wb') as f: f.write(await excel_template.read())
        ext = os.path.splitext(input_file.filename)[1].lower()
        if ext == '.pdf':
            data = extract_pdf_nuoc(ip) if rule == 'nuoc' else extract_pdf_dien(ip)
        elif ext in ('.xlsx', '.xls'): data = extract_excel_data(ip)
        else: raise HTTPException(400, f"Unsupported: {ext}")
        lk = build_lk(data)
        bn = os.path.basename(input_file.filename)
        m = re.search(r'(?:tháng|thang)\s*(\d+)[.\s]*(?:năm\s*)?(\d{4})?', bn, re.IGNORECASE)
        ms = int(m.group(1)) - RULES[rule]["ref_month"] if m else 0
        year = m.group(2) if m and m.group(2) else "2026"
        wb, res, u, sd, sn = process_phan_bo(ep, lk, rule, ms, year)
        mn = str(RULES[rule]["ref_month"] + ms).zfill(2)
        out_name = RULES[rule]["output_format"].format(month=mn, year=year) + ".xlsx"
        out_path = os.path.join(tmpdir, out_name)
        wb.save(out_path); wb.close()
        return FileResponse(out_path, filename=out_name,
                            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

@app.post("/doi-soat")
async def doi_soat(rule: str = Form(...), file1: UploadFile = File(...), file2: UploadFile = File(...)):
    if rule not in RULES: raise HTTPException(400, f"Invalid rule: {rule}")
    tmpdir = tempfile.mkdtemp()
    try:
        f1 = os.path.join(tmpdir, file1.filename)
        with open(f1, 'wb') as f: f.write(await file1.read())
        f2 = os.path.join(tmpdir, file2.filename)
        with open(f2, 'wb') as f: f.write(await file2.read())
        pgh_result, c1, c2 = process_doi_soat(f1, f2, rule)
        matched = sum(1 for v in pgh_result.values() if v == 'Đúng')
        mismatched = sum(1 for v in pgh_result.values() if v == 'Sai')
        no_result = sum(1 for v in pgh_result.values() if v == 'Không có kết quả')
        return {"rule": rule, "file1_count": c1, "file2_count": c2,
                "matched": matched, "mismatched": mismatched, "no_result": no_result, "total": len(pgh_result)}
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

@app.post("/doi-soat/download")
async def doi_soat_download(rule: str = Form(...), file1: UploadFile = File(...), file2: UploadFile = File(...)):
    if rule not in RULES: raise HTTPException(400, f"Invalid rule: {rule}")
    tmpdir = tempfile.mkdtemp()
    try:
        f1 = os.path.join(tmpdir, file1.filename)
        with open(f1, 'wb') as f: f.write(await file1.read())
        f2 = os.path.join(tmpdir, file2.filename)
        with open(f2, 'wb') as f: f.write(await file2.read())
        pgh_result, c1, c2 = process_doi_soat(f1, f2, rule)
        ext1 = os.path.splitext(f1)[1].lower()
        ext2 = os.path.splitext(f2)[1].lower()
        excel_file = f1 if ext1 in ('.xlsx','.xls') else (f2 if ext2 in ('.xlsx','.xls') else None)
        if excel_file:
            wb = openpyxl.load_workbook(excel_file, keep_vba=False)
            ws = wb.active
            rc = ws.max_column + 1
            ws.cell(row=1, column=rc).value = "Kết quả đối soát"
            ws.cell(row=1, column=rc).font = openpyxl.styles.Font(bold=True)
            for ri in range(2, ws.max_row + 1):
                pv = ws.cell(row=ri, column=5).value
                if pv:
                    m = re.search(r'\d{9}', str(pv))
                    if m:
                        result = pgh_result.get(m.group(0), 'Không có kết quả')
                        cell = ws.cell(row=ri, column=rc)
                        cell.value = result
                        cell.font = openpyxl.styles.Font(color='008000' if result=='Đúng' else ('FF0000' if result=='Sai' else '808080'), bold=True)
                        cell.alignment = openpyxl.styles.Alignment(horizontal='center')
            out_name = f"Doi soat {os.path.splitext(os.path.basename(excel_file))[0]}.xlsx"
        else:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Doi soat"
            ws.cell(row=1, column=1).value = "Số PGH"
            ws.cell(row=1, column=2).value = "Kết quả"
            for ri, (pgh, result) in enumerate(sorted(pgh_result.items()), 2):
                ws.cell(row=ri, column=1).value = pgh
                ws.cell(row=ri, column=2).value = result
            out_name = "Doi soat.xlsx"
        out_path = os.path.join(tmpdir, out_name)
        wb.save(out_path); wb.close()
        return FileResponse(out_path, filename=out_name,
                            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8080)
